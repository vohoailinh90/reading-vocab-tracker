#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""vocab_tool.py — cap nhat tu vung Anh/Nhat va theo doi sach doc.

Nguon su that la cac file CSV trong data/ (git merge duoc theo tung dong).
File Excel data/vocabulary.xlsx la ban SINH RA tu CSV — khong sua tay.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import subprocess
import sys
import time
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG_PATH = REPO_ROOT / "config" / "vocab.config.json"

LANGS = ("en", "ja")

VOCAB_FIELDS = [
    "id",
    "added_at",
    "term",
    "reading",
    "meaning_vi",
    "register",
    "spoken_equivalent",
    "example",
    "example_vi",
    "source",
    "page",
    "tags",
    "notes",
]

BOOK_FIELDS = [
    "id",
    "title",
    "author",
    "language",
    "status",
    "progress",
    "started_at",
    "finished_at",
    "rating",
    "notes_path",
    "updated_at",
]

# Nhan tieng Viet dung lam header trong file Excel. IMPORT_SYNONYMS ben duoi
# chua dung cac nhan nay nen xuat ra Excel roi import lai khong mat cot.
VOCAB_LABELS = {
    "id": "id",
    "added_at": "ngày thêm",
    "term": "từ vựng",
    "reading": "cách đọc",
    "meaning_vi": "nghĩa tiếng Việt",
    "register": "văn nói/văn viết",
    "spoken_equivalent": "cụm văn nói tương đương",
    "example": "ví dụ",
    "example_vi": "nghĩa ví dụ",
    "source": "nguồn",
    "page": "trang",
    "tags": "nhãn",
    "notes": "ghi chú",
}

BOOK_LABELS = {
    "id": "id",
    "title": "tên sách",
    "author": "tác giả",
    "language": "ngôn ngữ",
    "status": "trạng thái",
    "progress": "tiến độ",
    "started_at": "bắt đầu",
    "finished_at": "hoàn thành",
    "rating": "đánh giá",
    "notes_path": "ghi chú (file)",
    "updated_at": "cập nhật",
}

COLUMN_WIDTHS = {
    "id": 6,
    "added_at": 12,
    "term": 26,
    "reading": 22,
    "meaning_vi": 34,
    "register": 18,
    "spoken_equivalent": 26,
    "example": 44,
    "example_vi": 44,
    "source": 24,
    "page": 8,
    "tags": 18,
    "notes": 30,
    "title": 34,
    "author": 22,
    "language": 10,
    "status": 14,
    "progress": 12,
    "started_at": 12,
    "finished_at": 12,
    "rating": 10,
    "notes_path": 28,
    "updated_at": 12,
}

DEFAULT_CONFIG: Dict = {
    "data_dir": "data",
    "excel_path": "data/vocabulary.xlsx",
    # Duong dan file Excel cu o ngoai repo. De null cho toi khi biet duong dan.
    "external_excel_path": None,
    "git": {"remote": "origin", "branch": "main"},
}

# Hiragana, katakana, kanji, katakana nua rong.
JA_PATTERN = re.compile(r"[\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uff66-\uff9f]")

# Header (da chuan hoa) -> ten truong, dung khi import mot workbook la.
IMPORT_SYNONYMS: Dict[str, str] = {}


def _register_synonyms(field: str, *headers: str) -> None:
    for header in headers:
        IMPORT_SYNONYMS[_norm_header(header)] = field


def _norm_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _norm_key(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


class ToolError(Exception):
    """Loi do nguoi dung nhap sai — thoat voi ma 2, khong phai traceback."""


# ---------------------------------------------------------------- config / io


def load_config(path: Optional[Path] = None) -> Dict:
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    path = path or DEFAULT_CONFIG_PATH
    if path.exists():
        raw = json.loads(path.read_text(encoding="utf-8"))
        for key, value in raw.items():
            if key.startswith("_"):  # "_comment", "_todo" ... chi de ghi chu
                continue
            if isinstance(value, dict) and isinstance(cfg.get(key), dict):
                cfg[key].update(value)
            else:
                cfg[key] = value
    cfg["_config_path"] = str(path)
    return cfg


def resolve_path(value: Optional[str]) -> Optional[Path]:
    if not value:
        return None
    path = Path(str(value)).expanduser()
    return path if path.is_absolute() else (REPO_ROOT / path)


def vocab_csv(cfg: Dict, lang: str) -> Path:
    return resolve_path(cfg["data_dir"]) / f"vocabulary.{lang}.csv"


def books_csv(cfg: Dict) -> Path:
    return resolve_path(cfg["data_dir"]) / "books.csv"


def read_rows(path: Path, fields: List[str]) -> List[Dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8", newline="") as handle:
        return [
            {field: (row.get(field) or "").strip() for field in fields}
            for row in csv.DictReader(handle)
        ]


def write_rows(path: Path, fields: List[str], rows: List[Dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        # lineterminator="\n": mac dinh cua module csv la "\r\n", con
        # .gitattributes chuan hoa data/*.csv ve LF — de mac dinh thi moi lan
        # dung lai file se hien "da sua" du noi dung khong doi.
        writer = csv.DictWriter(
            handle, fieldnames=fields, extrasaction="ignore", lineterminator="\n"
        )
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def next_id(rows: List[Dict[str, str]]) -> int:
    used = [int(row["id"]) for row in rows if str(row.get("id", "")).isdigit()]
    return max(used) + 1 if used else 1


def detect_lang(term: str) -> str:
    return "ja" if JA_PATTERN.search(term or "") else "en"


def upsert(
    rows: List[Dict[str, str]],
    fields: List[str],
    key_field: str,
    values: Dict[str, str],
    strict: bool = False,
) -> Tuple[str, Dict[str, str]]:
    """Them moi hoac cap nhat theo key_field.

    Khi trung, chi ghi de nhung truong duoc truyen vao lan nay — cac o cu
    khong bi xoa trang.
    """
    key = _norm_key(values.get(key_field, ""))
    if not key:
        raise ToolError(f"thiếu giá trị bắt buộc: --{key_field}")

    for row in rows:
        if _norm_key(row.get(key_field, "")) == key:
            if strict:
                return ("duplicate", row)
            for field, value in values.items():
                if field in fields and value not in (None, ""):
                    row[field] = value
            return ("updated", row)

    row = {field: "" for field in fields}
    row.update({k: v for k, v in values.items() if k in fields and v is not None})
    row["id"] = str(next_id(rows))
    rows.append(row)
    return ("added", row)


# ------------------------------------------------------------------ excel out


def build_xlsx(cfg: Dict) -> List[Path]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - phu thuoc moi truong
        raise ToolError(
            "chưa cài openpyxl — chạy: python3 -m pip install -r requirements.txt"
        ) from exc

    sheets = [
        ("EN", VOCAB_FIELDS, VOCAB_LABELS, read_rows(vocab_csv(cfg, "en"), VOCAB_FIELDS)),
        ("JA", VOCAB_FIELDS, VOCAB_LABELS, read_rows(vocab_csv(cfg, "ja"), VOCAB_FIELDS)),
        ("Books", BOOK_FIELDS, BOOK_LABELS, read_rows(books_csv(cfg), BOOK_FIELDS)),
    ]

    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")

    for name, fields, labels, rows in sheets:
        sheet = workbook.create_sheet(name)
        sheet.append([labels[field] for field in fields])
        for row in rows:
            sheet.append([_cell_value(row.get(field, "")) for field in fields])

        for index, field in enumerate(fields, start=1):
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = COLUMN_WIDTHS.get(field, 18)
            cell = sheet.cell(row=1, column=index)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(fields))}{max(sheet.max_row, 1)}"
        )

    targets = [resolve_path(cfg["excel_path"])]
    external = resolve_path(cfg.get("external_excel_path"))
    if external:
        targets.append(external)

    written = []
    for target in targets:
        _save_deterministic(workbook, target)
        written.append(target)
    return written


# Moc thoi gian co dinh: khong mang y nghia ngay thang, chi de cung du lieu
# thi ra cung byte.
_FIXED_STAMP = datetime(2020, 1, 1)
_FIXED_ZIP_DATE = (1980, 1, 1, 0, 0, 0)


def _save_deterministic(workbook, target: Path) -> None:
    """Ghi .xlsx sao cho cung du lieu -> cung byte.

    Mac dinh openpyxl nhet thoi diem tao/sua vao docProps va thoi gian hien tai
    vao tung entry cua zip, nen dung lai cung mot du lieu van ra file khac byte.
    Trong repo dung chung thi moi lan `build-xlsx` se sinh ra mot commit binary
    rong va lam tang xung dot vo co.
    """
    workbook.properties.created = _FIXED_STAMP
    workbook.properties.modified = _FIXED_STAMP

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    target.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(
        target, "w", zipfile.ZIP_DEFLATED
    ) as dest:
        # Giu nguyen thu tu entry openpyxl sinh ra; chi chuan hoa moc thoi gian.
        for item in source.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=_FIXED_ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = item.external_attr
            payload = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                payload = _pin_modified_stamp(payload)
            dest.writestr(info, payload)


_MODIFIED_TAG = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _pin_modified_stamp(payload: bytes) -> bytes:
    """openpyxl ghi de dcterms:modified bang gio hien tai luc save, bat ke gia
    tri da dat trong workbook.properties — nen phai chuan hoa o buoc nay."""
    stamp = _FIXED_STAMP.strftime("%Y-%m-%dT%H:%M:%SZ").encode("ascii")
    return _MODIFIED_TAG.sub(rb"\g<1>" + stamp + rb"\g<2>", payload)


def _cell_value(value: str):
    text = (value or "").strip()
    if text.isdigit() and len(text) < 10:
        return int(text)
    return text


# -------------------------------------------------------------------- commands


def cmd_init(args, cfg) -> int:
    for lang in LANGS:
        path = vocab_csv(cfg, lang)
        if not path.exists():
            write_rows(path, VOCAB_FIELDS, [])
            print(f"đã tạo {path.relative_to(REPO_ROOT) if _under_root(path) else path}")
    path = books_csv(cfg)
    if not path.exists():
        write_rows(path, BOOK_FIELDS, [])
        print(f"đã tạo {path.relative_to(REPO_ROOT) if _under_root(path) else path}")
    for target in build_xlsx(cfg):
        print(f"đã dựng {target}")
    return 0


def cmd_add(args, cfg) -> int:
    lang = args.lang if args.lang != "auto" else detect_lang(args.term)
    path = vocab_csv(cfg, lang)
    rows = read_rows(path, VOCAB_FIELDS)

    values = {
        "added_at": args.date or date.today().isoformat(),
        "term": args.term.strip(),
        "reading": args.reading,
        "meaning_vi": args.meaning,
        "register": args.register,
        "spoken_equivalent": args.spoken,
        "example": args.example,
        "example_vi": args.example_vi,
        "source": args.source,
        "page": args.page,
        "tags": args.tags,
        "notes": args.notes,
    }

    action, row = upsert(rows, VOCAB_FIELDS, "term", values, strict=args.strict)
    if action == "duplicate":
        print(
            f"[{lang}] '{row['term']}' đã có (id={row['id']}). "
            "Bỏ --strict để cập nhật đè.",
            file=sys.stderr,
        )
        return 2

    write_rows(path, VOCAB_FIELDS, rows)
    label = "đã thêm" if action == "added" else "đã cập nhật"
    print(f"[{lang}] {label} id={row['id']}: {row['term']} = {row['meaning_vi']}")
    if not args.no_build:
        for target in build_xlsx(cfg):
            print(f"đã dựng {target}")
    return 0


def cmd_add_batch(args, cfg) -> int:
    payload = json.loads(Path(args.file).read_text(encoding="utf-8"))
    if not isinstance(payload, list):
        raise ToolError("file batch phải là một mảng JSON các object")

    counts = {"added": 0, "updated": 0, "skipped": 0}
    cache: Dict[str, List[Dict[str, str]]] = {}

    for index, entry in enumerate(payload, start=1):
        term = str(entry.get("term", "")).strip()
        if not term:
            print(f"bỏ qua mục {index}: thiếu 'term'", file=sys.stderr)
            counts["skipped"] += 1
            continue
        lang = entry.get("lang") or detect_lang(term)
        if lang not in LANGS:
            print(f"bỏ qua mục {index}: lang '{lang}' không hợp lệ", file=sys.stderr)
            counts["skipped"] += 1
            continue
        rows = cache.setdefault(lang, read_rows(vocab_csv(cfg, lang), VOCAB_FIELDS))
        values = {field: str(entry.get(field, "")).strip() for field in VOCAB_FIELDS}
        values["term"] = term
        values["added_at"] = values["added_at"] or date.today().isoformat()
        values.pop("id", None)
        action, _ = upsert(rows, VOCAB_FIELDS, "term", values)
        counts[action] += 1

    for lang, rows in cache.items():
        write_rows(vocab_csv(cfg, lang), VOCAB_FIELDS, rows)

    print(
        f"thêm {counts['added']}, cập nhật {counts['updated']}, bỏ qua {counts['skipped']}"
    )
    if not args.no_build:
        for target in build_xlsx(cfg):
            print(f"đã dựng {target}")
    return 0


def cmd_list(args, cfg) -> int:
    langs = LANGS if args.lang == "all" else (args.lang,)
    needle = _norm_key(args.grep) if args.grep else ""
    total = 0
    for lang in langs:
        rows = read_rows(vocab_csv(cfg, lang), VOCAB_FIELDS)
        if needle:
            rows = [
                row
                for row in rows
                if needle in _norm_key(" ".join(row.get(f, "") for f in VOCAB_FIELDS))
            ]
        rows = rows[-args.limit :] if args.limit else rows
        if not rows:
            continue
        print(f"--- {lang.upper()} ({len(rows)}) ---")
        for row in rows:
            suffix = f"  [{row['source']}]" if row.get("source") else ""
            reading = f" {row['reading']}" if row.get("reading") else ""
            print(f"  {row['id']:>4}  {row['term']}{reading} = {row['meaning_vi']}{suffix}")
        total += len(rows)
    if total == 0:
        print("chưa có từ nào khớp.")
    return 0


def cmd_book(args, cfg) -> int:
    path = books_csv(cfg)
    rows = read_rows(path, BOOK_FIELDS)
    values = {
        "title": args.title.strip(),
        "author": args.author,
        "language": args.language,
        "status": args.status,
        "progress": args.progress,
        "started_at": args.started,
        "finished_at": args.finished,
        "rating": args.rating,
        "notes_path": args.notes_path,
        "updated_at": date.today().isoformat(),
    }
    action, row = upsert(rows, BOOK_FIELDS, "title", values, strict=args.strict)
    if action == "duplicate":
        print(f"'{row['title']}' đã có (id={row['id']}).", file=sys.stderr)
        return 2
    write_rows(path, BOOK_FIELDS, rows)
    label = "đã thêm" if action == "added" else "đã cập nhật"
    print(f"sách: {label} id={row['id']}: {row['title']} — {row['status']}")
    if not args.no_build:
        for target in build_xlsx(cfg):
            print(f"đã dựng {target}")
    return 0


def cmd_build(args, cfg) -> int:
    for target in build_xlsx(cfg):
        print(f"đã dựng {target}")
    return 0


def cmd_import(args, cfg) -> int:
    """Nap mot file Excel co san (vi du file tu vung cu) vao CSV."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ToolError(
            "chưa cài openpyxl — chạy: python3 -m pip install -r requirements.txt"
        ) from exc

    source = Path(args.path).expanduser()
    if not source.exists():
        raise ToolError(f"không tìm thấy file: {source}")

    overrides = {}
    for pair in args.map or []:
        if "=" not in pair:
            raise ToolError(f"--map phải có dạng 'Tên cột=trường': {pair}")
        header, field = pair.split("=", 1)
        field = field.strip()
        if field not in VOCAB_FIELDS:
            raise ToolError(
                f"trường '{field}' không hợp lệ. Hợp lệ: {', '.join(VOCAB_FIELDS)}"
            )
        overrides[_norm_header(header)] = field

    # Khong dung read_only=True: can doc sheet.merged_cells de nhan ra dong tieu
    # de hoac ghi chu bi gop o — thu rat hay gap trong file Excel lam tay.
    workbook = load_workbook(source, data_only=True)
    sheet_names = args.sheet or workbook.sheetnames
    cache: Dict[str, List[Dict[str, str]]] = {}
    counts = {"added": 0, "updated": 0, "skipped": 0}

    for sheet_name in sheet_names:
        if sheet_name not in workbook.sheetnames:
            print(f"bỏ qua sheet '{sheet_name}': không tồn tại", file=sys.stderr)
            continue

        sheet = workbook[sheet_name]
        grid = list(sheet.iter_rows(values_only=True))
        if not grid:
            print(f"bỏ qua sheet '{sheet_name}': trống", file=sys.stderr)
            continue

        header_index, mapping, ignored = _find_header(grid, overrides, args.header_row)
        if mapping is None:
            print(
                f"bỏ qua sheet '{sheet_name}': không nhận ra cột từ vựng trong "
                f"{min(len(grid), _HEADER_SCAN_DEPTH)} dòng đầu. Dùng --header-row N "
                "để chỉ đúng dòng tiêu đề, hoặc --map 'Tên cột=term'.",
                file=sys.stderr,
            )
            continue

        for header, field in ignored:
            print(
                f"sheet '{sheet_name}': bỏ qua cột trùng '{header}' — đã có cột "
                f"khác nhận '{field}'",
                file=sys.stderr,
            )

        merged_rows = _horizontally_merged_rows(sheet)
        print(
            f"sheet '{sheet_name}' (tiêu đề ở dòng {header_index + 1}): "
            + ", ".join(sorted(set(mapping.values())))
        )

        for number, raw in enumerate(grid[header_index + 1 :], start=header_index + 2):
            if number in merged_rows:
                print(
                    f"  bỏ qua dòng {number}: ô gộp ngang, không phải dòng dữ liệu",
                    file=sys.stderr,
                )
                counts["skipped"] += 1
                continue

            entry = {
                field: str(raw[index]).strip()
                for index, field in mapping.items()
                if index < len(raw) and raw[index] not in (None, "")
            }
            term = entry.get("term", "")
            if not term:
                counts["skipped"] += 1
                continue

            lang = args.lang if args.lang != "auto" else detect_lang(term)
            entry.setdefault("added_at", date.today().isoformat())
            entry.setdefault("source", args.source or source.name)

            if args.dry_run:
                counts["added"] += 1
                print(f"  [{lang}] {term} = {entry.get('meaning_vi', '')}")
                continue

            rows = cache.setdefault(lang, read_rows(vocab_csv(cfg, lang), VOCAB_FIELDS))
            action, _ = upsert(rows, VOCAB_FIELDS, "term", entry)
            counts[action] += 1

    workbook.close()

    if args.dry_run:
        print(f"[thử] sẽ nạp {counts['added']} dòng, bỏ qua {counts['skipped']}")
        return 0

    for lang, rows in cache.items():
        write_rows(vocab_csv(cfg, lang), VOCAB_FIELDS, rows)
    print(
        f"thêm {counts['added']}, cập nhật {counts['updated']}, bỏ qua {counts['skipped']}"
    )
    if not args.no_build:
        for target in build_xlsx(cfg):
            print(f"đã dựng {target}")
    return 0


# So dong dau duoc quet de tim hang tieu de. File Excel lam tay thuong co mot
# vai dong tua de phia tren bang du lieu that.
_HEADER_SCAN_DEPTH = 10


def _find_header(grid, overrides, forced):
    """Tra ve (chi so dong tieu de, mapping, danh sach cot trung bi bo qua).

    mapping = None nghia la khong tim thay cot tu vung.
    """
    if forced:
        index = forced - 1
        if not 0 <= index < len(grid):
            raise ToolError(
                f"--header-row {forced} nằm ngoài phạm vi sheet ({len(grid)} dòng)"
            )
        mapping, ignored = _build_mapping(grid[index], overrides)
        return index, (mapping if "term" in mapping.values() else None), ignored

    for index, row in enumerate(grid[:_HEADER_SCAN_DEPTH]):
        if not any(cell not in (None, "") for cell in row):
            continue
        mapping, ignored = _build_mapping(row, overrides)
        if "term" in mapping.values():
            return index, mapping, ignored
    return 0, None, []


def _build_mapping(header_row, overrides):
    mapping: Dict[int, str] = {}
    ignored: List[Tuple[str, str]] = []
    for index, cell in enumerate(header_row):
        key = _norm_header(cell)
        if not key:
            continue
        field = overrides.get(key) or IMPORT_SYNONYMS.get(key)
        if not field or field == "id":
            continue
        if field in mapping.values():
            ignored.append((str(cell).strip(), field))
            continue
        mapping[index] = field
    return mapping, ignored


def _horizontally_merged_rows(sheet) -> set:
    """Cac dong nam trong vung gop o theo chieu ngang.

    Trong file lam tay, do gan nhu luon la dong tua de hoac ghi chu trai ngang
    bang — nap chung vao se tao ra nhung "tu vung" rac.
    """
    rows: set = set()
    ranges = getattr(getattr(sheet, "merged_cells", None), "ranges", []) or []
    for cell_range in ranges:
        if cell_range.max_col > cell_range.min_col:
            rows.update(range(cell_range.min_row, cell_range.max_row + 1))
    return rows


def cmd_sync(args, cfg) -> int:
    remote = cfg["git"].get("remote", "origin")
    branch = cfg["git"].get("branch", "main")
    paths = [
        path
        for path in (str(resolve_path(cfg["data_dir"])), str(REPO_ROOT / "books"))
        if Path(path).exists()
    ]
    if not paths:
        print("không tìm thấy thư mục dữ liệu để commit.", file=sys.stderr)
        return 2

    _git("add", "--", *paths)
    # Moi lenh git deu gioi han theo `paths`: neu nguoi dung da `git add` san mot
    # file khong lien quan, commit tu vung khong duoc cuon file do vao.
    staged = subprocess.run(
        ["git", "diff", "--cached", "--quiet", "--", *paths], cwd=REPO_ROOT, check=False
    )
    if staged.returncode == 0:
        print("không có thay đổi để commit.")
        return 0

    message = args.message or f"vocab: cập nhật từ vựng ({date.today().isoformat()})"
    _git("commit", "-m", message, "--", *paths)
    print(f"đã commit: {message}")

    if args.no_push:
        print("bỏ qua push (--no-push).")
        return 0

    delay = 2
    for attempt in range(1, 5):
        result = subprocess.run(
            ["git", "push", "-u", remote, branch],
            cwd=REPO_ROOT,
            check=False,
            capture_output=True,
            text=True,
        )
        if result.returncode == 0:
            print(f"đã push lên {remote}/{branch}.")
            return 0
        sys.stderr.write(result.stderr)
        if "rejected" in result.stderr or "non-fast-forward" in result.stderr:
            print(
                "push bị từ chối vì remote có commit mới. Chạy:\n"
                f"  git pull --no-rebase {remote} {branch}\n"
                "  # nếu data/vocabulary.xlsx xung đột: git checkout --ours "
                "data/vocabulary.xlsx && python3 tools/vocab_tool.py build-xlsx\n"
                f"  git push {remote} {branch}",
                file=sys.stderr,
            )
            return 2
        if attempt < 4:
            print(f"push lỗi mạng, thử lại sau {delay}s...", file=sys.stderr)
            time.sleep(delay)
            delay *= 2
    return 2


def _git(*args: str) -> None:
    subprocess.run(["git", *args], cwd=REPO_ROOT, check=True)


def _under_root(path: Path) -> bool:
    try:
        path.relative_to(REPO_ROOT)
        return True
    except ValueError:
        return False


# ----------------------------------------------------------------- arg parsing


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="vocab_tool.py",
        description="Cập nhật từ vựng Anh/Nhật ra Excel và theo dõi sách đọc.",
    )
    parser.add_argument("--config", help="đường dẫn file cấu hình JSON")
    subparsers = parser.add_subparsers(dest="command", required=True)

    init = subparsers.add_parser("init", help="tạo CSV rỗng và file Excel")
    init.set_defaults(func=cmd_init)

    add = subparsers.add_parser("add", help="thêm/cập nhật một từ")
    add.add_argument("--term", required=True, help="từ hoặc cụm từ")
    add.add_argument("--lang", choices=("auto", *LANGS), default="auto")
    add.add_argument("--reading", default="", help="IPA (Anh) hoặc kana (Nhật)")
    add.add_argument("--meaning", default="", help="nghĩa tiếng Việt")
    add.add_argument("--register", default="", help="văn nói / văn viết / trung tính")
    add.add_argument("--spoken", default="", help="cụm văn nói tương đương")
    add.add_argument("--example", default="", help="câu ví dụ")
    add.add_argument("--example-vi", dest="example_vi", default="", help="nghĩa câu ví dụ")
    add.add_argument("--source", default="", help="nguồn / tên sách")
    add.add_argument("--page", default="", help="trang")
    add.add_argument("--tags", default="", help="nhãn, cách nhau bằng dấu phẩy")
    add.add_argument("--notes", default="", help="ghi chú")
    add.add_argument("--date", default="", help="ngày thêm (mặc định hôm nay)")
    add.add_argument("--strict", action="store_true", help="báo lỗi nếu từ đã có")
    add.add_argument("--no-build", action="store_true", help="không dựng lại Excel")
    add.set_defaults(func=cmd_add)

    batch = subparsers.add_parser("add-batch", help="thêm nhiều từ từ file JSON")
    batch.add_argument("file", help="file JSON chứa mảng object")
    batch.add_argument("--no-build", action="store_true")
    batch.set_defaults(func=cmd_add_batch)

    listing = subparsers.add_parser("list", help="xem từ đã lưu")
    listing.add_argument("--lang", choices=("all", *LANGS), default="all")
    listing.add_argument("--limit", type=int, default=20, help="0 = tất cả")
    listing.add_argument("--grep", default="", help="lọc theo chuỗi")
    listing.set_defaults(func=cmd_list)

    book = subparsers.add_parser("book", help="thêm/cập nhật một cuốn sách")
    book.add_argument("--title", required=True)
    book.add_argument("--author", default="")
    book.add_argument("--language", default="")
    book.add_argument("--status", default="", help="chưa đọc / đang đọc / đã đọc / tạm dừng")
    book.add_argument("--progress", default="", help="ví dụ: 120/320 hoặc 40%%")
    book.add_argument("--started", default="")
    book.add_argument("--finished", default="")
    book.add_argument("--rating", default="")
    book.add_argument("--notes-path", dest="notes_path", default="", help="books/<ten>.md")
    book.add_argument("--strict", action="store_true")
    book.add_argument("--no-build", action="store_true")
    book.set_defaults(func=cmd_book)

    build = subparsers.add_parser("build-xlsx", help="dựng lại Excel từ CSV")
    build.set_defaults(func=cmd_build)

    importer = subparsers.add_parser("import", help="nạp một file Excel có sẵn vào CSV")
    importer.add_argument("--path", required=True, help="đường dẫn file .xlsx")
    importer.add_argument("--sheet", action="append", help="chỉ nạp sheet này (lặp được)")
    importer.add_argument("--lang", choices=("auto", *LANGS), default="auto")
    importer.add_argument(
        "--header-row",
        type=int,
        help="dòng chứa tiêu đề cột (1 là dòng đầu); mặc định tự dò 10 dòng đầu",
    )
    importer.add_argument("--source", default="", help="ghi đè cột nguồn")
    importer.add_argument(
        "--map",
        action="append",
        help="ánh xạ cột thủ công, ví dụ --map 'Từ vựng=term'",
    )
    importer.add_argument("--dry-run", action="store_true", help="chỉ in ra, không ghi")
    importer.add_argument("--no-build", action="store_true")
    importer.set_defaults(func=cmd_import)

    sync = subparsers.add_parser("sync", help="commit và push dữ liệu lên repo")
    sync.add_argument("-m", "--message", default="")
    sync.add_argument("--no-push", action="store_true")
    sync.set_defaults(func=cmd_sync)

    return parser


def main(argv: Optional[List[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    cfg = load_config(Path(args.config).expanduser() if args.config else None)
    try:
        return args.func(args, cfg)
    except ToolError as exc:
        print(f"lỗi: {exc}", file=sys.stderr)
        return 2


_register_synonyms(
    "term", "term", "word", "vocabulary", "expression", "phrase", "english", "japanese",
    "từ", "tu", "từ vựng", "tu vung", "cụm từ", "cum tu", "単語", "語",
)
_register_synonyms(
    "reading", "reading", "ipa", "pronunciation", "phonetic", "furigana", "kana",
    "hiragana", "katakana", "romaji", "cách đọc", "cach doc", "phát âm", "phat am",
    "phiên âm", "phien am", "読み", "よみ", "ふりがな",
)
_register_synonyms(
    "meaning_vi", "meaning", "translation", "vietnamese", "nghĩa", "nghia",
    "nghĩa tiếng việt", "nghia tieng viet", "ý nghĩa", "y nghia", "dịch", "dich",
    "意味",
)
_register_synonyms(
    "register", "register", "style", "formality", "văn nói/văn viết",
    "van noi/van viet", "sắc thái", "sac thai", "ngữ vực", "ngu vuc",
)
_register_synonyms(
    "spoken_equivalent", "spoken", "spoken equivalent", "cụm văn nói tương đương",
    "cum van noi tuong duong", "văn nói tương đương", "van noi tuong duong",
)
_register_synonyms(
    "example", "example", "sentence", "example sentence", "ví dụ", "vi du",
    "câu ví dụ", "cau vi du", "例文", "例",
)
_register_synonyms(
    "example_vi", "example_vi", "example meaning", "nghĩa ví dụ", "nghia vi du",
    "dịch ví dụ", "dich vi du", "ví dụ tiếng việt", "vi du tieng viet",
)
_register_synonyms(
    "source", "source", "book", "reference", "nguồn", "nguon", "sách", "sach",
    "tài liệu", "tai lieu", "xuất xứ", "xuat xu",
)
_register_synonyms("page", "page", "trang", "số trang", "so trang", "ページ")
_register_synonyms(
    "tags", "tags", "tag", "category", "topic", "nhãn", "nhan", "chủ đề", "chu de",
    "phân loại", "phan loai",
)
_register_synonyms(
    "notes", "notes", "note", "remark", "ghi chú", "ghi chu", "chú thích", "chu thich",
)
_register_synonyms(
    "added_at", "date", "added", "added_at", "ngày", "ngay", "ngày thêm", "ngay them",
    "ngày học", "ngay hoc",
)

if __name__ == "__main__":  # pragma: no cover
    sys.exit(main())
