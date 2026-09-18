# -*- coding: utf-8 -*-
"""Kiem thu cho tools/vocab_tool.py."""

from __future__ import annotations

import csv
import json
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))

import vocab_tool  # noqa: E402


@pytest.fixture()
def workspace(tmp_path: Path) -> Path:
    """Config tro toan bo duong dan vao tmp_path, khong dung cham repo that."""
    config = {
        "data_dir": str(tmp_path / "data"),
        "excel_path": str(tmp_path / "data" / "vocabulary.xlsx"),
        "external_excel_path": None,
    }
    config_path = tmp_path / "config.json"
    config_path.write_text(json.dumps(config), encoding="utf-8")
    return config_path


def run(config_path: Path, *args: str) -> int:
    return vocab_tool.main(["--config", str(config_path), *args])


def read_csv(config_path: Path, name: str) -> list[dict]:
    data_dir = Path(json.loads(config_path.read_text(encoding="utf-8"))["data_dir"])
    path = data_dir / name
    if not path.exists():
        return []
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def test_add_writes_row_and_builds_excel(workspace: Path) -> None:
    assert run(workspace, "add", "--term", "hostage negotiator",
               "--meaning", "người đàm phán con tin", "--lang", "en") == 0

    rows = read_csv(workspace, "vocabulary.en.csv")
    assert len(rows) == 1
    assert rows[0]["term"] == "hostage negotiator"
    assert rows[0]["meaning_vi"] == "người đàm phán con tin"
    assert rows[0]["id"] == "1"

    xlsx = Path(json.loads(workspace.read_text(encoding="utf-8"))["excel_path"])
    sheet = load_workbook(xlsx)["EN"]
    assert sheet.cell(row=1, column=3).value == "từ vựng"
    assert sheet.cell(row=2, column=3).value == "hostage negotiator"


def test_duplicate_term_updates_instead_of_appending(workspace: Path) -> None:
    run(workspace, "add", "--term", "root cause", "--meaning", "nguyên nhân gốc",
        "--source", "Bosch DNOx report")
    run(workspace, "add", "--term", "  Root Cause  ", "--meaning", "nguyên nhân cốt lõi")

    rows = read_csv(workspace, "vocabulary.en.csv")
    assert len(rows) == 1
    assert rows[0]["meaning_vi"] == "nguyên nhân cốt lõi"
    # Truong khong truyen lai lan hai phai duoc giu nguyen, khong bi xoa trang.
    assert rows[0]["source"] == "Bosch DNOx report"


def test_strict_refuses_duplicate(workspace: Path) -> None:
    run(workspace, "add", "--term", "deadline", "--meaning", "hạn chót")
    assert run(workspace, "add", "--term", "deadline", "--meaning", "khác",
               "--strict") == 2
    rows = read_csv(workspace, "vocabulary.en.csv")
    assert len(rows) == 1
    assert rows[0]["meaning_vi"] == "hạn chót"


@pytest.mark.parametrize(
    ("term", "expected"),
    [("納期", "ja"), ("しめきり", "ja"), ("ナットク", "ja"), ("deadline", "en"),
     ("root cause analysis", "en")],
)
def test_language_autodetection(term: str, expected: str) -> None:
    assert vocab_tool.detect_lang(term) == expected


def test_autodetect_routes_japanese_to_ja_sheet(workspace: Path) -> None:
    run(workspace, "add", "--term", "納期", "--reading", "のうき", "--meaning", "kỳ hạn giao hàng")
    assert read_csv(workspace, "vocabulary.en.csv") == []
    rows = read_csv(workspace, "vocabulary.ja.csv")
    assert len(rows) == 1 and rows[0]["reading"] == "のうき"


def test_add_batch_reads_json_array(workspace: Path, tmp_path: Path) -> None:
    batch = tmp_path / "batch.json"
    batch.write_text(
        json.dumps(
            [
                {"term": "calibration", "meaning_vi": "hiệu chuẩn"},
                {"term": "残業", "meaning_vi": "tăng ca"},
                {"meaning_vi": "thiếu term nên bị bỏ qua"},
            ],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    assert run(workspace, "add-batch", str(batch)) == 0
    assert len(read_csv(workspace, "vocabulary.en.csv")) == 1
    assert len(read_csv(workspace, "vocabulary.ja.csv")) == 1


def test_import_maps_vietnamese_headers(workspace: Path, tmp_path: Path) -> None:
    legacy = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tu vung"
    sheet.append(["Từ vựng", "Cách đọc", "Nghĩa tiếng Việt", "Ví dụ"])
    sheet.append(["escalate", "/ˈeskəleɪt/", "leo thang / đẩy lên cấp trên", "We escalate it."])
    sheet.append(["点検", "てんけん", "kiểm tra định kỳ", "点検を行う。"])
    workbook.save(legacy)

    assert run(workspace, "import", "--path", str(legacy)) == 0

    en = read_csv(workspace, "vocabulary.en.csv")
    ja = read_csv(workspace, "vocabulary.ja.csv")
    assert [row["term"] for row in en] == ["escalate"]
    assert en[0]["meaning_vi"] == "leo thang / đẩy lên cấp trên"
    assert en[0]["reading"] == "/ˈeskəleɪt/"
    assert en[0]["source"] == "legacy.xlsx"
    assert [row["term"] for row in ja] == ["点検"]


def test_import_dry_run_writes_nothing(workspace: Path, tmp_path: Path) -> None:
    legacy = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    workbook.active.append(["Từ vựng", "Nghĩa"])
    workbook.active.append(["tolerance", "dung sai"])
    workbook.save(legacy)

    assert run(workspace, "import", "--path", str(legacy), "--dry-run") == 0
    assert read_csv(workspace, "vocabulary.en.csv") == []


def test_import_custom_column_mapping(workspace: Path, tmp_path: Path) -> None:
    legacy = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    workbook.active.append(["Cot la", "Dien giai"])
    workbook.active.append(["throughput", "thông lượng"])
    workbook.save(legacy)

    # Khong co --map thi khong nhan ra cot tu vung -> khong nap gi.
    assert run(workspace, "import", "--path", str(legacy)) == 0
    assert read_csv(workspace, "vocabulary.en.csv") == []

    assert run(workspace, "import", "--path", str(legacy),
               "--map", "Cot la=term", "--map", "Dien giai=meaning_vi") == 0
    rows = read_csv(workspace, "vocabulary.en.csv")
    assert rows[0]["term"] == "throughput" and rows[0]["meaning_vi"] == "thông lượng"


def test_excel_roundtrip_survives_reimport(workspace: Path, tmp_path: Path) -> None:
    """Excel do tool sinh ra phai import lai duoc — nhan tieng Viet khop synonyms."""
    run(workspace, "add", "--term", "tailpipe", "--meaning", "ống xả", "--reading", "/ˈteɪlpaɪp/")
    generated = Path(json.loads(workspace.read_text(encoding="utf-8"))["excel_path"])

    second = tmp_path / "second.json"
    second.write_text(
        json.dumps({"data_dir": str(tmp_path / "data2"),
                    "excel_path": str(tmp_path / "data2" / "vocabulary.xlsx")}),
        encoding="utf-8",
    )
    assert run(second, "import", "--path", str(generated), "--sheet", "EN") == 0

    rows = read_csv(second, "vocabulary.en.csv")
    assert len(rows) == 1
    assert rows[0]["term"] == "tailpipe"
    assert rows[0]["meaning_vi"] == "ống xả"
    assert rows[0]["reading"] == "/ˈteɪlpaɪp/"


def test_book_add_then_update(workspace: Path) -> None:
    run(workspace, "book", "--title", "Never Split the Difference",
        "--author", "Chris Voss", "--status", "đang đọc", "--progress", "40/300")
    run(workspace, "book", "--title", "never split the difference", "--status", "đã đọc")

    rows = read_csv(workspace, "books.csv")
    assert len(rows) == 1
    assert rows[0]["status"] == "đã đọc"
    assert rows[0]["author"] == "Chris Voss"
    assert rows[0]["progress"] == "40/300"


def test_init_creates_all_files(workspace: Path) -> None:
    assert run(workspace, "init") == 0
    config = json.loads(workspace.read_text(encoding="utf-8"))
    data_dir = Path(config["data_dir"])
    for name in ("vocabulary.en.csv", "vocabulary.ja.csv", "books.csv"):
        assert (data_dir / name).exists()
    assert load_workbook(config["excel_path"]).sheetnames == ["EN", "JA", "Books"]


def test_ids_keep_increasing(workspace: Path) -> None:
    for term in ("alpha", "beta", "gamma"):
        run(workspace, "add", "--term", term, "--meaning", term, "--no-build")
    rows = read_csv(workspace, "vocabulary.en.csv")
    assert [row["id"] for row in rows] == ["1", "2", "3"]


def test_csv_uses_lf_line_endings(workspace: Path) -> None:
    """CSV phải dùng LF: .gitattributes chuẩn hoá data/*.csv về LF, nếu tool ghi
    CRLF thì mỗi lần dựng lại file sẽ hiện 'đã sửa' dù nội dung không đổi."""
    run(workspace, "add", "--term", "drift", "--meaning", "trôi/lệch dần", "--no-build")
    data_dir = Path(json.loads(workspace.read_text(encoding="utf-8"))["data_dir"])
    raw = (data_dir / "vocabulary.en.csv").read_bytes()
    assert b"\r\n" not in raw
    assert raw.count(b"\n") == 2  # header + 1 dòng


def test_xlsx_build_is_byte_stable(workspace: Path) -> None:
    """Cùng dữ liệu phải ra cùng byte, nếu không mỗi lần dựng lại là một commit
    binary rỗng và một xung đột merge vô cớ trong repo dùng chung."""
    import zipfile

    run(workspace, "add", "--term", "torque", "--meaning", "mô-men xoắn", "--no-build")
    xlsx = Path(json.loads(workspace.read_text(encoding="utf-8"))["excel_path"])

    run(workspace, "build-xlsx")
    first = xlsx.read_bytes()
    with zipfile.ZipFile(xlsx) as archive:
        stamps = {item.date_time for item in archive.infolist()}

    run(workspace, "build-xlsx")
    assert xlsx.read_bytes() == first
    assert stamps == {(1980, 1, 1, 0, 0, 0)}

    # Hai lan dung o tren co the roi vao cung mot giay, nen kiem tra thang moc
    # thoi gian trong docProps: day moi la thu openpyxl ghi de theo gio he thong.
    with zipfile.ZipFile(xlsx) as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
    assert core.count("2020-01-01T00:00:00Z") == 2
    assert str(date.today().year) not in core


def test_xlsx_stays_readable_after_normalising(workspace: Path) -> None:
    run(workspace, "add", "--term", "backlash", "--meaning", "độ rơ / phản ứng dữ dội")
    xlsx = Path(json.loads(workspace.read_text(encoding="utf-8"))["excel_path"])
    book = load_workbook(xlsx)
    assert book.sheetnames == ["EN", "JA", "Books"]
    assert book["EN"].cell(row=2, column=3).value == "backlash"


def _legacy_workbook(path: Path, *, title_row: bool = False, note_row: bool = False,
                     duplicate_header: bool = False) -> Path:
    """Dựng một file Excel 'làm tay' giống file từ vựng cũ của người dùng."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tu vung cu"
    if title_row:
        sheet.append(["TỪ VỰNG BOSCH DNOx - 2024", None, None])
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    header = ["Từ vựng", "Nghĩa tiếng Việt", "Ví dụ"]
    if duplicate_header:
        header = ["Từ vựng", "Nghĩa tiếng Việt", "Nghĩa tiếng Việt"]
    sheet.append(header)
    sheet.append(["tailpipe", "ống xả", "Check the tailpipe."])
    if note_row:
        note_index = sheet.max_row + 1
        sheet.append(["Ghi chú chung cho cả nhóm từ bên dưới", None, None])
        sheet.merge_cells(start_row=note_index, start_column=1, end_row=note_index,
                          end_column=3)
    sheet.append(["catalyst", "chất xúc tác", "The catalyst is degraded."])
    workbook.save(path)
    return path


def test_import_skips_merged_title_row_above_header(workspace: Path, tmp_path: Path) -> None:
    """Dòng tựa đề gộp ô phía trên bảng không được nhận nhầm là hàng tiêu đề."""
    legacy = _legacy_workbook(tmp_path / "title.xlsx", title_row=True)
    assert run(workspace, "import", "--path", str(legacy)) == 0

    rows = read_csv(workspace, "vocabulary.en.csv")
    assert [row["term"] for row in rows] == ["tailpipe", "catalyst"]
    assert rows[0]["meaning_vi"] == "ống xả"


def test_import_skips_merged_note_row_inside_table(workspace: Path, tmp_path: Path) -> None:
    """Dòng ghi chú gộp ngang giữa bảng không được nạp thành một 'từ vựng' rác."""
    legacy = _legacy_workbook(tmp_path / "note.xlsx", note_row=True)
    assert run(workspace, "import", "--path", str(legacy)) == 0

    terms = [row["term"] for row in read_csv(workspace, "vocabulary.en.csv")]
    assert terms == ["tailpipe", "catalyst"]
    assert not any("Ghi chú" in term for term in terms)


def test_import_warns_on_duplicate_header_columns(
    workspace: Path, tmp_path: Path, capsys: pytest.CaptureFixture
) -> None:
    """Cột tiêu đề trùng tên bị bỏ, nhưng phải báo ra chứ không im lặng."""
    legacy = _legacy_workbook(tmp_path / "dup.xlsx", duplicate_header=True)
    assert run(workspace, "import", "--path", str(legacy)) == 0

    captured = capsys.readouterr()
    assert "bỏ qua cột trùng" in captured.err
    assert "meaning_vi" in captured.err


def test_import_header_row_override(workspace: Path, tmp_path: Path) -> None:
    """--header-row chỉ đúng dòng tiêu đề khi phần tựa đề dài quá 10 dòng."""
    legacy = tmp_path / "deep.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for index in range(12):
        sheet.append([f"dòng tựa đề {index}"])
    sheet.append(["Từ vựng", "Nghĩa tiếng Việt"])
    sheet.append(["throughput", "thông lượng"])
    workbook.save(legacy)

    # Tu do chi quet 10 dong dau nen khong thay tieu de.
    assert run(workspace, "import", "--path", str(legacy)) == 0
    assert read_csv(workspace, "vocabulary.en.csv") == []

    assert run(workspace, "import", "--path", str(legacy), "--header-row", "13") == 0
    rows = read_csv(workspace, "vocabulary.en.csv")
    assert [row["term"] for row in rows] == ["throughput"]


def test_import_header_row_out_of_range(workspace: Path, tmp_path: Path) -> None:
    legacy = _legacy_workbook(tmp_path / "small.xlsx")
    assert run(workspace, "import", "--path", str(legacy), "--header-row", "99") == 2


def test_import_dry_run_still_writes_nothing_with_title_row(
    workspace: Path, tmp_path: Path
) -> None:
    legacy = _legacy_workbook(tmp_path / "dry.xlsx", title_row=True, note_row=True)
    assert run(workspace, "import", "--path", str(legacy), "--dry-run") == 0
    assert read_csv(workspace, "vocabulary.en.csv") == []


@pytest.fixture()
def git_repo(tmp_path: Path) -> Path:
    """Một repo git thật với remote bare, để chạy `sync` end-to-end.

    Tool suy ra gốc repo từ vị trí file của chính nó, nên chỉ cần chép
    vocab_tool.py vào <tmp>/repo/tools/ là gốc repo thành <tmp>/repo.
    """
    import shutil
    import subprocess

    if shutil.which("git") is None:  # pragma: no cover
        pytest.skip("không có git")

    repo = tmp_path / "repo"
    (repo / "tools").mkdir(parents=True)
    (repo / "config").mkdir()
    shutil.copy(Path(vocab_tool.__file__), repo / "tools" / "vocab_tool.py")
    (repo / "config" / "vocab.config.json").write_text(
        json.dumps({"data_dir": "data", "excel_path": "data/vocabulary.xlsx",
                    "git": {"remote": "origin", "branch": "main"}}),
        encoding="utf-8",
    )

    remote = tmp_path / "remote.git"
    subprocess.run(["git", "init", "-q", "--bare", "-b", "main", str(remote)], check=True)
    for command in (
        ["git", "init", "-q", "-b", "main"],
        ["git", "config", "user.name", "Test"],
        ["git", "config", "user.email", "t@example.com"],
        ["git", "remote", "add", "origin", str(remote)],
        ["git", "add", "-A"],
        ["git", "commit", "-q", "-m", "init"],
        ["git", "push", "-q", "-u", "origin", "main"],
    ):
        subprocess.run(command, cwd=repo, check=True)
    return repo


def _run_tool(repo: Path, *args: str):
    import subprocess

    return subprocess.run(
        [sys.executable, "tools/vocab_tool.py", *args],
        cwd=repo, capture_output=True, text=True,
    )


def test_sync_commits_only_data_paths(git_repo: Path) -> None:
    """Nếu người dùng đã `git add` một file không liên quan, commit từ vựng
    không được cuốn file đó vào."""
    import subprocess

    _run_tool(git_repo, "init")
    _run_tool(git_repo, "add", "--term", "urea", "--meaning", "dung dịch urê", "--no-build")

    (git_repo / "README.md").write_text("sửa tay", encoding="utf-8")
    subprocess.run(["git", "add", "README.md"], cwd=git_repo, check=True)

    result = _run_tool(git_repo, "sync", "-m", "vocab: chỉ dữ liệu")
    assert result.returncode == 0, result.stderr

    changed = subprocess.run(
        ["git", "show", "--name-only", "--format=", "HEAD"],
        cwd=git_repo, capture_output=True, text=True, check=True,
    ).stdout.split()
    assert "README.md" not in changed
    assert "data/vocabulary.en.csv" in changed


def test_sync_reports_nothing_to_commit(git_repo: Path) -> None:
    _run_tool(git_repo, "init")
    _run_tool(git_repo, "sync", "-m", "lần đầu")
    result = _run_tool(git_repo, "sync", "-m", "lần hai")
    assert result.returncode == 0
    assert "không có thay đổi" in result.stdout


def test_sync_pushes_to_remote(git_repo: Path) -> None:
    import subprocess

    _run_tool(git_repo, "init")
    _run_tool(git_repo, "add", "--term", "soot", "--meaning", "muội than", "--no-build")
    assert _run_tool(git_repo, "sync", "-m", "vocab: muội than").returncode == 0

    remote = git_repo.parent / "remote.git"
    log = subprocess.run(
        ["git", "--git-dir", str(remote), "log", "--oneline", "-1", "main"],
        capture_output=True, text=True, check=True,
    ).stdout
    assert "vocab: muội than" in log
